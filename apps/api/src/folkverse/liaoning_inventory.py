"""Operator-only inventory intake. Metadata collection never publishes chatbot evidence."""

import argparse
import hashlib
import io
import json
import re
import subprocess
import tempfile
from datetime import UTC, datetime
from html.parser import HTMLParser
from pathlib import Path
from typing import Any, Literal
from urllib.parse import urlsplit
from urllib.robotparser import RobotFileParser

import httpx
import openpyxl  # type: ignore[import-untyped]
import xlrd  # type: ignore[import-untyped]
from pydantic import BaseModel, ConfigDict, Field, model_validator

from folkverse.config import ROOT

VERSION = "liaoning-inventory-v1"
EXTRACTOR_VERSION = "liaoning-metadata-extractor-v4"
MAX_BYTES = 32 * 1024 * 1024
AGENT = "FolkVerseInventory/1.0"
SUBJECTS = [
    "history_geography",
    "folklore",
    "belief",
    "festivals",
    "crafts",
    "artifacts",
    "performance",
    "food",
    "clothing",
    "museums",
    "sites",
    "people",
    "terminology",
]


class SourceSpec(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)
    id: str = Field(pattern=r"^[a-z0-9-]{1,80}$")
    title: str
    institution: str
    url: str
    resource_url: str
    adapter: Literal[
        "workbook",
        "html_table",
        "html_lines",
        "html_codes",
        "html_article",
        "pdf_table",
        "api_json",
        "descriptor",
    ]
    kind: Literal[
        "geography", "heritage", "museum", "sites", "history", "checkpoint", "explanation"
    ]
    reference_date: str
    edition: str
    expected_count: int | None = Field(default=None, ge=0)
    counting_unit: str
    sheet: str | None = None
    rights: str
    check_interval_days: int = Field(default=365, ge=1)
    notes: str = ""
    locality: str = ""
    subject_hint: str | None = None


def digest(value: Any) -> str:
    return hashlib.sha256(
        json.dumps(value, sort_keys=True, ensure_ascii=False).encode()
    ).hexdigest()


def write_json(path: Path, value: Any) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(value, ensure_ascii=False, indent=2) + "\n")
    temporary.replace(path)


def registry(folder: Path) -> list[SourceSpec]:
    document = json.loads((folder / "registry.json").read_text())
    sources = [SourceSpec.model_validate(s) for s in document["sources"]]
    if len({s.id for s in sources}) != len(sources):
        raise ValueError("Duplicate source identity")
    return sources


def safe_url(url: str) -> str:
    parsed = urlsplit(url)
    if (
        parsed.scheme != "https"
        or parsed.username
        or parsed.password
        or parsed.port not in {None, 443}
        or not parsed.hostname
        or not any(
            parsed.hostname == h or parsed.hostname.endswith("." + h)
            for h in ("ln.gov.cn", "mct.gov.cn")
        )
    ):
        raise ValueError("Inventory access requires a registered official HTTPS host")
    return url


def download(url: str, client: httpx.Client) -> bytes:
    safe_url(url)
    host = urlsplit(url)
    robots_url = f"https://{host.netloc}/robots.txt"
    response = client.get(robots_url)
    if response.status_code == 200:
        robots = RobotFileParser()
        robots.parse(response.text.splitlines())
        if not robots.can_fetch(AGENT, url):
            raise ValueError("robots_disallowed")
    elif response.status_code not in {404, 410}:
        raise ValueError(f"robots_unavailable_http_{response.status_code}")
    with client.stream("GET", url) as response:
        if response.status_code != 200:
            raise ValueError(f"resource_http_{response.status_code}")
        data = bytearray()
        for chunk in response.iter_bytes():
            if len(data) + len(chunk) > MAX_BYTES:
                raise ValueError("resource_too_large")
            data.extend(chunk)
    return bytes(data)


class Tables(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.tables: list[list[list[str]]] = []
        self.lines: list[str] = []
        self.table: list[list[str]] | None = None
        self.row: list[str] | None = None
        self.cell: list[str] | None = None
        self.skip = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style"}:
            self.skip += 1
        if tag == "table":
            self.table = []
        elif tag == "tr" and self.table is not None:
            self.row = []
        elif tag in {"td", "th"} and self.row is not None:
            self.cell = []
        if tag in {"p", "br", "tr", "div"}:
            self.lines.append("\n")

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"}:
            self.skip = max(0, self.skip - 1)
        if tag in {"td", "th"} and self.cell is not None and self.row is not None:
            self.row.append("".join(self.cell).strip())
            self.cell = None
        elif tag == "tr" and self.row is not None and self.table is not None:
            self.table.append(self.row)
            self.row = None
        elif tag == "table" and self.table is not None:
            self.tables.append(self.table)
            self.table = None
        if tag in {"p", "tr", "div"}:
            self.lines.append("\n")

    def handle_data(self, data: str) -> None:
        if self.skip:
            return
        self.lines.append(data)
        if self.cell is not None:
            self.cell.append(data)


class Article(HTMLParser):
    """Extract only the department's article container, excluding site chrome/scripts."""

    def __init__(self) -> None:
        super().__init__()
        self.depth = 0
        self.active: int | None = None
        self.current: list[str] | None = None
        self.paragraphs: list[str] = []
        self.skip = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if tag in {"script", "style"}:
            self.skip += 1
        if tag == "div":
            self.depth += 1
            if "pages_content" in (dict(attrs).get("class") or "").split():
                self.active = self.depth
        if tag == "p" and self.active is not None:
            self.current = []

    def handle_endtag(self, tag: str) -> None:
        if tag in {"script", "style"}:
            self.skip = max(0, self.skip - 1)
        if tag == "p" and self.current is not None:
            self.paragraphs.append("".join(self.current).strip())
            self.current = None
        if tag == "div":
            if self.active == self.depth:
                self.active = None
            self.depth -= 1

    def handle_data(self, data: str) -> None:
        if self.current is not None and not self.skip:
            self.current.append(data)


ROMAN_CATEGORIES = {
    "Ⅰ": "民间文学",
    "Ⅱ": "传统音乐",
    "Ⅲ": "传统舞蹈",
    "Ⅳ": "传统戏剧",
    "Ⅴ": "曲艺",
    "Ⅵ": "传统体育、游艺与杂技",
    "Ⅶ": "传统美术",
    "Ⅷ": "传统技艺",
    "Ⅸ": "传统医药",
    "Ⅹ": "民俗",
}


def clean(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def workbook(raw: bytes, sheet: str | None) -> list[tuple[str, list[list[str]]]]:
    result = []
    if raw.startswith(b"PK"):
        import zipfile

        with zipfile.ZipFile(io.BytesIO(raw)) as zipped:
            if sum(f.file_size for f in zipped.infolist()) > 64 * 1024 * 1024:
                raise ValueError("Expanded workbook exceeds limit")
        book = openpyxl.load_workbook(io.BytesIO(raw), data_only=True)
        for page in book:
            if sheet is not None and page.title != sheet:
                continue
            rows = [[clean(v) for v in values] for values in page.values]
            for span in page.merged_cells.ranges:
                value = clean(page.cell(span.min_row, span.min_col).value)
                for y in range(span.min_row - 1, span.max_row):
                    for x in range(span.min_col - 1, span.max_col):
                        rows[y][x] = value
            result.append((page.title, rows))
        book.close()
    else:
        legacy = xlrd.open_workbook(file_contents=raw)
        for page in legacy.sheets():
            if sheet is not None and page.name != sheet:
                continue
            rows = [[clean(v) for v in page.row_values(i)] for i in range(page.nrows)]
            for y0, y1, x0, x1 in page.merged_cells:
                value = rows[y0][x0]
                for y in range(y0, y1):
                    for x in range(x0, x1):
                        rows[y][x] = value
            result.append((page.name, rows))
    if sheet is not None and not result:
        raise ValueError("Registered worksheet is absent")
    return result


def record(
    source: SourceSpec,
    ordinal: str,
    name: str,
    category: str,
    locality: str,
    locator: str,
    native_code: str = "",
    **extra: Any,
) -> dict[str, Any]:
    identifier = "row_" + digest([source.id, ordinal, locator])[:24]
    return {
        "id": identifier,
        "ordinal": ordinal,
        "name_original": name,
        "category_original": category,
        "locality_original": locality,
        "native_code": native_code,
        "locator": locator,
        "entity_id": "entity_" + digest([source.id, ordinal])[:24],
        "disposition": "represented" if name else "unresolved",
        "errors": [],
        **extra,
    }


def extract(source: SourceSpec, raw: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    rows: list[dict[str, Any]] = []
    errors: list[str] = []
    if source.adapter == "html_article":
        article = Article()
        article.feed(raw.decode("utf-8"))
        paragraphs = [
            {
                "locator": f"article:p:{number}",
                "text_sha256": hashlib.sha256(text.encode()).hexdigest(),
                "characters": len(text),
            }
            for number, text in enumerate(article.paragraphs, 1)
            if len(text) > 40
        ]
        if not paragraphs:
            return [], ["No substantive article paragraphs found in registered container"]
        return [
            record(
                source,
                "article",
                source.title,
                source.subject_hint or "",
                source.locality,
                "article:pages_content",
                paragraphs=paragraphs,
                source_text_storage="private_original_capture_only",
                content_approval="pending_human_review",
            )
        ], []
    if source.adapter == "descriptor":
        return rows, ["Descriptor/checkpoint only; no enumeration imported"]
    if source.adapter == "workbook":
        for sheet, values in workbook(raw, source.sheet):
            if source.kind == "geography":
                found: dict[str, dict[str, Any]] = {}
                for number, cells in enumerate(values, 1):
                    for column, level in [(1, "city"), (2, "county")]:
                        if len(cells) <= column:
                            continue
                        match = re.search(r"([^\s]+)\s+(21\d{4})(?!\d)", cells[column])
                        if not match:
                            continue
                        name, code = match.groups()
                        if code == "210000":
                            continue
                        found.setdefault(
                            code,
                            record(
                                source,
                                code,
                                name,
                                level,
                                "",
                                f"sheet:{sheet};row:{number};column:{column + 1}",
                                code,
                                level=level,
                                city_code=code[:4] + "00" if level == "county" else code,
                            ),
                        )
                rows.extend(found.values())
            elif source.kind == "museum":
                for number, cells in enumerate(values, 1):
                    if len(cells) >= 4 and re.fullmatch(r"\d+", cells[0]):
                        rows.append(
                            record(
                                source,
                                cells[0],
                                cells[1],
                                cells[3],
                                cells[2],
                                f"sheet:{sheet};row:{number}",
                                repository_address=cells[6] if len(cells) > 6 else "",
                            )
                        )
            else:
                category = ""
                for number, cells in enumerate(values, 1):
                    if len(cells) >= 3 and re.fullmatch(r"\d+", cells[0]):
                        category = cells[1] or category
                        rows.append(
                            record(
                                source,
                                cells[0],
                                cells[2],
                                category,
                                "",
                                f"sheet:{sheet};row:{number}",
                            )
                        )
    elif source.adapter in {"html_table", "html_lines", "html_codes"}:
        parser = Tables()
        parser.feed(raw.decode("utf-8"))
        category = ""
        if source.adapter == "html_table":
            for table_number, table in enumerate(parser.tables, 1):
                for row_number, cells in enumerate(table, 1):
                    if len(cells) == 1 and re.search(r"[一二三四五六七八九十]+、", cells[0]):
                        category = cells[0]
                    if (
                        source.kind == "sites"
                        and len(cells) == 5
                        and re.fullmatch(r"\d+", cells[0].strip())
                    ):
                        rows.append(
                            record(
                                source,
                                cells[0],
                                cells[1],
                                cells[4],
                                cells[2],
                                f"table:{table_number};row:{row_number}",
                                chronology_original=cells[3],
                            )
                        )
                    elif len(cells) == 4 and re.fullmatch(r"\d+", cells[0].strip()):
                        rows.append(
                            record(
                                source,
                                cells[0],
                                cells[2],
                                category or ROMAN_CATEGORIES.get(cells[1][:1], ""),
                                cells[3],
                                f"table:{table_number};row:{row_number}",
                                cells[1],
                            )
                        )
        elif source.adapter == "html_codes":
            for number, line in enumerate("".join(parser.lines).splitlines(), 1):
                match = re.search(r"(?:^|\s)(\d{1,3})\s+([ⅠⅡⅢⅣⅤⅥⅦⅧⅨⅩ][-—－]\d+)\s+(.+)", line)
                if not match:
                    continue
                ordinal, code, remaining = match.groups()
                parts = re.split(r"\s+", remaining.strip())
                if len(parts) < 2:
                    errors.append(f"Unsplit native row {ordinal}")
                    continue
                rows.append(
                    record(
                        source,
                        ordinal,
                        " ".join(parts[:-1]),
                        ROMAN_CATEGORIES.get(code[:1], ""),
                        parts[-1],
                        f"text-line:{number}",
                        code,
                    )
                )
        else:
            for number, line in enumerate("".join(parser.lines).splitlines(), 1):
                line = line.strip()
                if re.match(r"[一二三四五六七八九十]+、", line):
                    category = line
                match = re.fullmatch(r"(\d{1,3})\s+(.+?)\s+([^\s]+)", line)
                if match:
                    ordinal, name, locality = match.groups()
                    rows.append(
                        record(
                            source,
                            ordinal,
                            name,
                            category,
                            locality.rstrip("\ue004"),
                            f"text-line:{number}",
                        )
                    )
    elif source.adapter == "api_json":
        document = json.loads(raw)
        for number, entry in enumerate(document["items"], 1):
            if (
                not isinstance(entry, dict)
                or not {"id", "name", "category", "locality"} <= entry.keys()
                or not all(isinstance(entry[key], str) for key in ("name", "category", "locality"))
                or not isinstance(entry["id"], (str, int))
                or not str(entry["id"]).strip()
            ):
                errors.append(f"Invalid API inventory item {number}")
                continue
            rows.append(
                record(
                    source,
                    str(entry["id"]),
                    str(entry["name"]),
                    str(entry["category"]),
                    str(entry["locality"]),
                    f"items:{number}",
                )
            )
    else:
        with tempfile.TemporaryDirectory() as temporary:
            pdf, text = Path(temporary) / "input.pdf", Path(temporary) / "text.txt"
            pdf.write_bytes(raw)
            subprocess.run(
                ["pdftotext", "-layout", str(pdf), str(text)],
                check=True,
                timeout=60,
                capture_output=True,
            )
            content = text.read_text()
            if not content.strip():
                return [], ["Image-only PDF; OCR and table reconciliation required"]
            for page_number, page in enumerate(content.split("\f"), 1):
                for line_number, line in enumerate(page.splitlines(), 1):
                    match = re.match(r"^\s*(\d+)\s+(.+?)\s{2,}(.+)$", line)
                    if match:
                        ordinal, name, locality = match.groups()
                        rows.append(
                            record(
                                source,
                                ordinal,
                                name,
                                "",
                                locality,
                                f"page:{page_number};line:{line_number}",
                            )
                        )
            errors.append(
                "PDF draft needs continuation/column reconciliation; no completeness claim"
            )
    if source.adapter == "pdf_table":
        for row in rows:
            row["disposition"] = "unresolved"
            row["errors"] = ["Unreconciled PDF columns"]
    if any(row["disposition"] == "unresolved" for row in rows):
        errors.append("Unresolved inventory rows require reconciliation")
    if not rows:
        errors.append("No inventory rows extracted")
    ordinals = [str(int(r["ordinal"])) if r["ordinal"].isdigit() else r["ordinal"] for r in rows]
    if len(ordinals) != len(set(ordinals)):
        errors.append("Duplicate native ordinals require explicit reconciliation")
    if source.expected_count is not None and len(set(ordinals)) != source.expected_count:
        errors.append(
            f"Expected {source.expected_count} {source.counting_unit}; "
            f"extracted {len(set(ordinals))}"
        )
    if source.kind in {"heritage", "museum"} and source.expected_count is not None:
        missing = sorted(
            set(map(str, range(1, source.expected_count + 1))) - set(ordinals), key=int
        )
        unexpected = sorted(set(ordinals) - set(map(str, range(1, source.expected_count + 1))))
        if missing or unexpected:
            errors.append(
                f"Native ordinal reconciliation: missing={missing}; unexpected={unexpected}"
            )
    return rows, errors


def ingest(folder: Path, source: SourceSpec, raw: bytes, fetched_at: str, raw_folder: Path) -> str:
    sha = hashlib.sha256(raw).hexdigest()
    key = digest([source.model_dump(), sha, EXTRACTOR_VERSION])
    destination = folder / "collections" / f"{source.id}-{key[:16]}.json"
    index_path = folder / "collection-index.json"
    index: dict[str, Any] = (
        json.loads(index_path.read_text())
        if index_path.exists()
        else {"version": VERSION, "sources": {}}
    )
    active = next((c for c in collections(folder) if c["source"]["id"] == source.id), None)
    if (
        active
        and active["source"] == source.model_dump()
        and active["raw_sha256"] == sha
        and active.get("extractor_version")
        in {"human-visual-transcription-v1", "assistant-visual-transcription-v1"}
    ):
        return str(active["collection_hash"])
    if destination.exists():
        saved = json.loads(destination.read_text())
        expected_rows, expected_errors = extract(source, raw)
        if (
            saved["source"] != source.model_dump()
            or saved["raw_sha256"] != sha
            or saved["collection_hash"] != key
            or saved["rows"] != expected_rows
            or saved["errors"] != expected_errors
        ):
            raise ValueError("Existing collection was modified; cannot bless altered intake")
    else:
        rows, errors = extract(source, raw)
        raw_folder.mkdir(parents=True, exist_ok=True)
        (raw_folder / sha).write_bytes(raw)
        write_json(
            destination,
            {
                "version": VERSION,
                "collection_hash": key,
                "source": source.model_dump(),
                "raw_sha256": sha,
                "original_filename": Path(urlsplit(source.resource_url).path).name,
                "fetched_at": fetched_at,
                "extractor_version": EXTRACTOR_VERSION,
                "status": "imported" if rows and not errors else "extraction_pending",
                "rows": rows,
                "errors": errors,
            },
        )
    index["sources"][source.id] = {
        "collection": destination.name,
        "collection_hash": key,
        "manifest_sha256": digest(json.loads(destination.read_text())),
    }
    write_json(index_path, index)
    return key


def collect(folder: Path, raw_folder: Path, selected: str | None = None) -> dict[str, Any]:
    known = {s.id for s in registry(folder)}
    if selected and selected not in known:
        raise ValueError("Unknown source identifier")
    previous = folder / "last-collection-run.json"
    attempts = (
        [
            a
            for a in json.loads(previous.read_text())["attempts"]
            if selected and a["source_id"] != selected
        ]
        if previous.exists()
        else []
    )
    with httpx.Client(timeout=25, headers={"User-Agent": AGENT}, follow_redirects=False) as client:
        for source in registry(folder):
            if selected and source.id != selected:
                continue
            try:
                raw = download(source.resource_url, client)
                key = ingest(folder, source, raw, datetime.now(UTC).isoformat(), raw_folder)
                attempts.append(
                    {"source_id": source.id, "result": "collected", "collection_hash": key}
                )
            except (httpx.HTTPError, ValueError, OSError, subprocess.SubprocessError) as error:
                # Do not copy response bodies, credentials or exception URLs into reports.
                reason = str(error) if isinstance(error, ValueError) else type(error).__name__
                attempts.append(
                    {"source_id": source.id, "result": "blocked", "reason": reason[:200]}
                )
    report = {"version": VERSION, "checked_at": datetime.now(UTC).isoformat(), "attempts": attempts}
    write_json(folder / "last-collection-run.json", report)
    return report


def collections(folder: Path) -> list[dict[str, Any]]:
    index_path = folder / "collection-index.json"
    if not index_path.exists():
        return []
    index = json.loads(index_path.read_text())
    result = []
    for entry in index["sources"].values():
        name = entry["collection"]
        if Path(name).name != name:
            raise ValueError("Invalid collection path")
        value = json.loads((folder / "collections" / name).read_text())
        if value["collection_hash"] != entry["collection_hash"] or digest(value) != entry.get(
            "manifest_sha256"
        ):
            raise ValueError("Collection pointer/hash mismatch")
        result.append(value)
    return result


class TranscribedRow(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)
    ordinal: int = Field(ge=1)
    name_original: str = Field(min_length=1)
    category_original: str
    locality_original: str
    native_code: str = ""
    chronology_original: str = ""
    extraction_notes: str = ""
    page_locator: str
    visual_row_locator: str = Field(min_length=1)


class Reconciliation(BaseModel):
    model_config = ConfigDict(extra="forbid", strict=True)
    source_id: str
    base_collection_hash: str
    observation_hash: str
    reviewer: str = Field(min_length=1)
    reviewer_kind: Literal["human", "assistant"] = "human"
    reviewed_at: str
    attest_visual_row_check: bool
    rows: list[TranscribedRow] = Field(min_length=1)
    notes: str = Field(min_length=1)

    @model_validator(mode="after")
    def attributed_check(self) -> "Reconciliation":
        if (
            not self.attest_visual_row_check
            or not self.reviewer.strip()
            or datetime.fromisoformat(self.reviewed_at).tzinfo is None
        ):
            raise ValueError("Transcription requires actual attributed visual checking")
        if len({r.ordinal for r in self.rows}) != len(self.rows):
            raise ValueError("Duplicate transcribed ordinals")
        if any(not row.name_original.strip() for row in self.rows):
            raise ValueError("Transcribed names must be nonblank")
        return self


def reconcile_scan(folder: Path, packet: Reconciliation) -> str:
    base = next(c for c in collections(folder) if c["source"]["id"] == packet.source_id)
    if base.get("reconciliation") == packet.model_dump():
        return str(base["collection_hash"])
    if base["collection_hash"] != packet.base_collection_hash:
        raise ValueError("Transcription is stale after inventory revision")
    if base["source"]["expected_count"] is None:
        raise ValueError("Scan inventory requires an explicit denominator")
    observation = next(
        (
            json.loads(p.read_text())
            for p in (folder / "ocr").glob("*.json")
            if json.loads(p.read_text())["observation_hash"] == packet.observation_hash
        ),
        None,
    )
    previous = base.get("reconciliation", {})
    same_capture_revision = (
        observation is not None
        and previous.get("observation_hash") == packet.observation_hash
        and observation.get("source_raw_sha256") == base["raw_sha256"]
    )
    if observation is None or (
        observation["source_collection_hash"] != base["collection_hash"]
        and not same_capture_revision
    ):
        raise ValueError("OCR observation does not match current inventory revision")
    # Observations themselves are immutable and checksum-bound.
    evidence = {k: v for k, v in observation.items() if k not in {"observation_hash", "created_at"}}
    if digest(evidence) != packet.observation_hash:
        raise ValueError("OCR observation checksum mismatch")
    pages = {p["locator"] for p in observation["pages"]}
    if any(r.page_locator not in pages for r in packet.rows):
        raise ValueError("Transcription locator missing from captured scans")
    source = SourceSpec.model_validate(base["source"])
    assert source.expected_count is not None
    if any(r.ordinal > source.expected_count for r in packet.rows):
        raise ValueError("Transcribed ordinal exceeds inventory denominator")
    key = digest([base["collection_hash"], packet.model_dump()])
    name = f"{source.id}-{key[:16]}.json"
    rows = [
        record(
            source,
            str(r.ordinal),
            r.name_original,
            r.category_original,
            r.locality_original,
            r.page_locator + ";" + r.visual_row_locator,
            r.native_code,
            chronology_original=r.chronology_original,
            extraction_notes=r.extraction_notes,
            extraction_review_kind=packet.reviewer_kind,
        )
        for r in packet.rows
    ]
    missing = sorted(set(range(1, source.expected_count + 1)) - {r.ordinal for r in packet.rows})
    value = {
        **base,
        "collection_hash": key,
        "rows": rows,
        "extractor_version": packet.reviewer_kind + "-visual-transcription-v1",
        "content_approval": "pending_human_review",
        "reconciliation": packet.model_dump(),
        "status": "imported" if not missing else "extraction_pending",
        "errors": [f"Missing native ordinals: {missing}"] if missing else [],
    }
    path = folder / "collections" / name
    if not path.exists():
        write_json(path, value)
    index_path = folder / "collection-index.json"
    index = json.loads(index_path.read_text())
    index["sources"][source.id] = {
        "collection": name,
        "collection_hash": key,
        "manifest_sha256": digest(value),
    }
    write_json(index_path, index)
    return key


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--folder", type=Path, default=ROOT / "data/liaoning")
    sub = parser.add_subparsers(dest="command", required=True)
    fetch = sub.add_parser("collect")
    fetch.add_argument("--source")
    local = sub.add_parser("import-file")
    local.add_argument("--source", required=True)
    local.add_argument("--file", type=Path, required=True)
    sub.add_parser("report")
    reconciliation = sub.add_parser("reconcile-scan")
    reconciliation.add_argument("--file", type=Path, required=True)
    args = parser.parse_args()
    raw_folder = ROOT / ".local/liaoning/raw"
    if args.command == "collect":
        result = collect(args.folder, raw_folder, args.source)
        print(json.dumps(result, ensure_ascii=False, indent=2))
    elif args.command == "reconcile-scan":
        print(
            reconcile_scan(
                args.folder, Reconciliation.model_validate(json.loads(args.file.read_text()))
            )
        )
    elif args.command == "import-file":
        source = next(s for s in registry(args.folder) if s.id == args.source)
        if args.file.stat().st_size > MAX_BYTES:
            raise ValueError("Local inventory exceeds limit")
        print(
            ingest(
                args.folder,
                source,
                args.file.read_bytes(),
                datetime.now(UTC).isoformat(),
                raw_folder,
            )
        )
    else:
        from folkverse.liaoning_review import coverage_report

        report = coverage_report(args.folder)
        print(json.dumps(report["summary"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
